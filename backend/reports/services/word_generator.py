"""
Word Document Generator Service
Generates formatted Word documents from extracted surveillance data
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from typing import Dict, Any, List
from datetime import datetime
import os
import tempfile


class WordReportGenerator:
    """Generate Word report from extracted KPI data"""

    # Path to Word template (relative to backend directory)
    TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'templates', 'report_template.docx')

    # Fallback: Path to logo image
    LOGO_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'assets', 'antic_logo.png')

    def __init__(self):
        self.doc = None
        self.using_template = False
        self._init_document()

    def _init_document(self):
        """Initialize document - use template if available, otherwise create new"""
        if os.path.exists(self.TEMPLATE_PATH):
            try:
                self.doc = Document(self.TEMPLATE_PATH)
                self.using_template = True
                print(f"Using Word template: {self.TEMPLATE_PATH}")
            except Exception as e:
                print(f"Error loading template: {e}, creating new document")
                self.doc = Document()
                self.using_template = False
        else:
            print(f"Template not found at {self.TEMPLATE_PATH}, creating new document")
            self.doc = Document()
            self.using_template = False

        self._setup_styles()
        self._add_page_numbers()

    def _setup_styles(self):
        """Configure document styles"""
        # Set default font
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)

    def _add_page_numbers(self):
        """Add page numbers to footer"""
        section = self.doc.sections[0]
        footer = section.footer
        footer.is_linked_to_previous = False

        # Add page number paragraph
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Add "Page X of Y" format
        run = paragraph.add_run("Page ")
        run.font.size = Pt(10)

        # Add PAGE field
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')

        instrText = OxmlElement('w:instrText')
        instrText.text = "PAGE"

        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')

        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)

        run2 = paragraph.add_run(" sur ")
        run2.font.size = Pt(10)

        # Add NUMPAGES field
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'begin')

        instrText2 = OxmlElement('w:instrText')
        instrText2.text = "NUMPAGES"

        fldChar4 = OxmlElement('w:fldChar')
        fldChar4.set(qn('w:fldCharType'), 'end')

        run2._r.append(fldChar3)
        run2._r.append(instrText2)
        run2._r.append(fldChar4)

    def generate_report(self, data: Dict[str, Any], output_path: str):
        """Generate complete Word report based on report type"""
        # Store generation time and report type
        self.generation_time = datetime.now()
        self.report_type = data.get('report_type', 'unknown')
        self.is_merged_report = 'source_files' in data and len(data.get('source_files', [])) > 1

        # Add cover page first (from template or fallback)
        self._add_cover_page(data)

        # Add page break after cover
        self.doc.add_page_break()

        # Add title
        self._add_title("RAPPORT DE VEILLE INFORMATIONNELLE")

        # For merged reports, show source files info
        if self.is_merged_report:
            source_files = data.get('source_files', [])
            self._add_subtitle(f"Rapport consolidé à partir de {len(source_files)} sources")
            for sf in source_files:
                self._add_source_file_info(sf)
        else:
            # Single source file info
            source_filename = data.get('source_filename', '')
            if source_filename:
                self._add_subtitle(f"Source: {source_filename}")

            # Add report type
            report_type_display = data.get('report_type_display', '')
            if report_type_display:
                self._add_subtitle(f"Type de rapport: {report_type_display}")

        # Add period
        period = data.get('period', {})
        if period.get('start') and period.get('end'):
            self._add_subtitle(f"Période: {period['start']} - {period['end']}")

        # Add generation timestamp
        self._add_generation_info()

        self.doc.add_paragraph()  # Spacing

        # Generate sections based on report type
        if self.is_merged_report:
            # Generate comprehensive merged report
            self._generate_merged_report(data)
        elif self.report_type == 'brand24_demographics':
            self._generate_demographics_report(data)
        elif self.report_type == 'brand24_analysis':
            self._generate_brand24_analysis_report(data)
        else:
            # Default: Mention.com report or unknown
            self._generate_mention_report(data)

        # Save document
        self.doc.save(output_path)
        return output_path

    def _add_source_file_info(self, source_info: Dict):
        """Add individual source file info line"""
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(f"• {source_info.get('filename', '')} ({source_info.get('report_type_display', source_info.get('report_type', ''))})")
        run.font.size = Pt(10)
        run.font.italic = True

    def _generate_merged_report(self, data: Dict[str, Any]):
        """Generate comprehensive report from merged data from multiple sources"""
        section_num = 1

        # 1. KPI Classification Section
        self._add_section_heading(f"{section_num}. CLASSIFICATION DES INDICATEURS CLÉS DE PERFORMANCE")

        # 1.1 Indicateurs de présence passive
        if self._has_data(data.get('presence_passive', {})):
            self._add_subsection_heading(f"{section_num}.1. Indicateurs de présence passive")
            self._add_presence_passive_table(data)

        # 1.2 Indicateurs de présence active
        if self._has_data(data.get('presence_active', {})):
            self._add_subsection_heading(f"{section_num}.2. Indicateurs de présence active")
            self._add_presence_active_table(data)

        # 1.3 Indicateurs des tendances d'opinions (Sentiment)
        if self._has_data(data.get('sentiment', {})):
            self._add_subsection_heading(f"{section_num}.3. Indicateurs des tendances d'opinions (Sentiment)")
            self._add_opinion_trends_table(data.get('sentiment', {}))

        # 1.4 Indicateurs des émotions
        if self._has_data(data.get('emotion', {})):
            self._add_subsection_heading(f"{section_num}.4. Indicateurs des émotions")
            self._add_emotions_table(data.get('emotion', {}))

        section_num += 1

        # Sources distribution
        if self._has_data(data.get('sources', {})):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR SOURCE")
            self._add_sources_table(data.get('sources', {}))
            section_num += 1

        # Languages
        if self._has_data(data.get('languages', {})):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR LANGUE")
            self._add_languages_table(data.get('languages', {}))
            section_num += 1

        # Topics
        if self._has_data(data.get('topics', [])):
            self._add_section_heading(f"{section_num}. SUJETS PRINCIPAUX")
            self._add_topics_table(data.get('topics', []))
            section_num += 1

        # Hashtags
        if self._has_data(data.get('hashtags', [])):
            self._add_section_heading(f"{section_num}. HASHTAGS POPULAIRES")
            self._add_hashtags_table(data.get('hashtags', []))
            section_num += 1

        # Influencers
        if self._has_data(data.get('influencers', [])):
            self._add_section_heading(f"{section_num}. INFLUENCEURS PRINCIPAUX")
            self._add_influencers_table(data.get('influencers', []))
            section_num += 1

        # Reach breakdown
        if self._has_data(data.get('reach_breakdown', [])):
            self._add_section_heading(f"{section_num}. PORTÉE DES PUBLICATIONS")
            self._add_reach_breakdown_table(data.get('reach_breakdown', []))
            section_num += 1

        # Demographics (if available from Brand24 Demographics report)
        demographics = data.get('demographics', {})
        if self._has_data(demographics):
            self._add_section_heading(f"{section_num}. DONNÉES DÉMOGRAPHIQUES")

            if self._has_data(demographics.get('gender', {})):
                self._add_subsection_heading(f"{section_num}.1. Répartition par genre")
                self._add_gender_table(demographics.get('gender', {}))

            if self._has_data(demographics.get('age', [])):
                self._add_subsection_heading(f"{section_num}.2. Répartition par âge")
                self._add_age_table(demographics.get('age', []))

            if self._has_data(demographics.get('countries', [])):
                self._add_subsection_heading(f"{section_num}.3. Répartition par pays")
                self._add_countries_table(demographics.get('countries', []))

            section_num += 1

        # Data Sources Reference section (for merged reports)
        data_sources = data.get('data_sources', {})
        if data_sources:
            self._add_section_heading(f"{section_num}. RÉFÉRENCES DES SOURCES DE DONNÉES")
            self._add_data_sources_reference(data_sources, data.get('source_files', []))

    def _add_data_sources_reference(self, data_sources: Dict, source_files: List[Dict]):
        """Add reference section showing which data came from which source"""
        # Add explanation
        para = self.doc.add_paragraph()
        run = para.add_run("Cette section indique la provenance des données utilisées dans ce rapport consolidé.")
        run.font.size = Pt(11)
        run.font.italic = True

        self.doc.add_paragraph()

        # Group data sources by source file
        sources_by_file = {}
        for field, sources in data_sources.items():
            for source in sources:
                filename = source.get('filename', 'Unknown')
                if filename not in sources_by_file:
                    sources_by_file[filename] = []
                # Extract the field category (first part before dot)
                field_category = field.split('.')[0] if '.' in field else field
                if field_category not in [f['category'] for f in sources_by_file[filename]]:
                    sources_by_file[filename].append({
                        'category': field_category,
                        'report_type': source.get('report_type', '')
                    })

        # Category labels in French
        category_labels = {
            'period_start': 'Période de début',
            'period_end': 'Période de fin',
            'presence_passive': 'Présence passive (followers, vues, portée)',
            'presence_active': 'Présence active (likes, commentaires, partages)',
            'volume': 'Volume de mentions',
            'reach': 'Portée',
            'sentiment': 'Sentiment',
            'emotion': 'Émotions',
            'sources': 'Sources/Plateformes',
            'languages': 'Langues',
            'topics': 'Sujets',
            'hashtags': 'Hashtags',
            'influencers': 'Influenceurs',
            'reach_breakdown': 'Portée des publications',
            'demographics': 'Données démographiques'
        }

        # Create table with source files and their data contributions
        table = self.doc.add_table(rows=len(sources_by_file) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Fichier Source'
        header_cells[1].text = 'Type de Rapport'
        header_cells[2].text = 'Données Fournies'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        # Data rows
        for idx, (filename, categories) in enumerate(sources_by_file.items(), 1):
            cells = table.rows[idx].cells
            cells[0].text = filename

            # Get report type
            report_type = categories[0].get('report_type', '') if categories else ''
            report_type_display = {
                'mention_dashboard': 'Mention.com Dashboard',
                'brand24_analysis': 'Brand24 Analysis',
                'brand24_demographics': 'Brand24 Demographics'
            }.get(report_type, report_type)
            cells[1].text = report_type_display

            # List categories
            category_names = [category_labels.get(c['category'], c['category']) for c in categories]
            cells[2].text = ', '.join(category_names[:5])  # Limit to 5 for readability
            if len(category_names) > 5:
                cells[2].text += f" (+{len(category_names) - 5} autres)"

    def _generate_mention_report(self, data: Dict[str, Any]):
        """Generate sections for Mention.com report"""
        section_num = 1

        # 1. KPI Classification Section
        self._add_section_heading(f"{section_num}. CLASSIFICATION DES INDICATEURS CLÉS DE PERFORMANCE")

        # 1.1 Indicateurs de présence passive
        if self._has_data(data.get('presence_passive', {})):
            self._add_subsection_heading(f"{section_num}.1. Indicateurs de présence passive")
            self._add_presence_passive_table(data)

        # 1.2 Indicateurs des tendances d'opinions (Sentiment)
        if self._has_data(data.get('sentiment', {})):
            self._add_subsection_heading(f"{section_num}.2. Indicateurs des tendances d'opinions (Sentiment)")
            self._add_opinion_trends_table(data.get('sentiment', {}))

        # 1.3 Indicateurs des émotions
        if self._has_data(data.get('emotion', {})):
            self._add_subsection_heading(f"{section_num}.3. Indicateurs des émotions")
            self._add_emotions_table(data.get('emotion', {}))

        section_num += 1

        # Sources distribution
        if self._has_data(data.get('sources', {})):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR SOURCE")
            self._add_sources_table(data.get('sources', {}))
            section_num += 1

        # Languages
        if self._has_data(data.get('languages', {})):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR LANGUE")
            self._add_languages_table(data.get('languages', {}))
            section_num += 1

        # Topics
        if self._has_data(data.get('topics', [])):
            self._add_section_heading(f"{section_num}. SUJETS PRINCIPAUX")
            self._add_topics_table(data.get('topics', []))
            section_num += 1

        # Hashtags
        if self._has_data(data.get('hashtags', [])):
            self._add_section_heading(f"{section_num}. HASHTAGS POPULAIRES")
            self._add_hashtags_table(data.get('hashtags', []))
            section_num += 1

        # Influencers
        if self._has_data(data.get('influencers', [])):
            self._add_section_heading(f"{section_num}. INFLUENCEURS PRINCIPAUX")
            self._add_influencers_table(data.get('influencers', []))
            section_num += 1

        # Reach breakdown (top posts)
        if self._has_data(data.get('reach_breakdown', [])):
            self._add_section_heading(f"{section_num}. PORTÉE DES PUBLICATIONS")
            self._add_reach_breakdown_table(data.get('reach_breakdown', []))

    def _generate_brand24_analysis_report(self, data: Dict[str, Any]):
        """Generate sections for Brand24 Analysis report"""
        section_num = 1

        # 1. KPI Classification Section
        self._add_section_heading(f"{section_num}. CLASSIFICATION DES INDICATEURS CLÉS DE PERFORMANCE")

        # 1.1 Indicateurs de présence passive
        if self._has_data(data.get('presence_passive', {})):
            self._add_subsection_heading(f"{section_num}.1. Indicateurs de présence passive")
            self._add_presence_passive_table(data)

        # 1.2 Indicateurs de présence active (likes, comments, shares)
        if self._has_data(data.get('presence_active', {})):
            self._add_subsection_heading(f"{section_num}.2. Indicateurs de présence active")
            self._add_presence_active_table(data)

        # 1.3 Indicateurs des tendances d'opinions (Sentiment)
        if self._has_data(data.get('sentiment', {})):
            self._add_subsection_heading(f"{section_num}.3. Indicateurs des tendances d'opinions (Sentiment)")
            self._add_opinion_trends_table(data.get('sentiment', {}))

        # 1.4 Presence Score (Brand24 specific)
        presence_score = data.get('presence_score', {})
        if self._has_data(presence_score):
            self._add_subsection_heading(f"{section_num}.4. Score de Présence")
            self._add_presence_score_info(presence_score)

        section_num += 1

        # Sources distribution
        if self._has_data(data.get('sources', {})):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR SOURCE")
            self._add_sources_table(data.get('sources', {}))
            section_num += 1

        # Hashtags
        if self._has_data(data.get('hashtags', [])):
            self._add_section_heading(f"{section_num}. HASHTAGS POPULAIRES")
            self._add_hashtags_table(data.get('hashtags', []))
            section_num += 1

        # Influencers
        if self._has_data(data.get('influencers', [])):
            self._add_section_heading(f"{section_num}. INFLUENCEURS PRINCIPAUX")
            self._add_influencers_table(data.get('influencers', []))
            section_num += 1

        # AVE (Advertising Value Equivalent)
        ave = data.get('ave', {})
        if self._has_data(ave):
            self._add_section_heading(f"{section_num}. VALEUR PUBLICITAIRE ÉQUIVALENTE (AVE)")
            self._add_ave_info(ave)

    def _generate_demographics_report(self, data: Dict[str, Any]):
        """Generate sections for Brand24 Demographics report"""
        section_num = 1
        demographics = data.get('demographics', {})

        # 1. Reach Overview
        reach = data.get('reach', {})
        if self._has_data(reach):
            self._add_section_heading(f"{section_num}. PORTÉE TOTALE")
            self._add_reach_summary(reach)
            section_num += 1

        # 2. Gender Distribution
        gender = demographics.get('gender', {})
        if self._has_data(gender):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR GENRE")
            self._add_gender_table(gender)
            section_num += 1

        # 3. Age Distribution
        age = demographics.get('age', [])
        if self._has_data(age):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR ÂGE")
            self._add_age_table(age)
            section_num += 1

        # 4. Countries Distribution
        countries = demographics.get('countries', [])
        if self._has_data(countries):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR PAYS")
            self._add_countries_table(countries)
            section_num += 1

        # 5. Occupation Distribution
        occupation = demographics.get('occupation', [])
        if self._has_data(occupation):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR OCCUPATION")
            self._add_occupation_table(occupation)
            section_num += 1

        # 6. Education Distribution
        education = demographics.get('education', [])
        if self._has_data(education):
            self._add_section_heading(f"{section_num}. RÉPARTITION PAR NIVEAU D'ÉDUCATION")
            self._add_education_table(education)
            section_num += 1

        # 7. Interests
        interests = demographics.get('interests', [])
        if self._has_data(interests):
            self._add_section_heading(f"{section_num}. CENTRES D'INTÉRÊT")
            self._add_interests_table(interests)

    def _add_generation_info(self):
        """Add report generation timestamp"""
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(f"Rapport généré le {self.generation_time.strftime('%d/%m/%Y à %H:%M')}")
        run.font.size = Pt(10)
        run.font.italic = True
        run.font.color.rgb = RGBColor(128, 128, 128)

    def _add_cover_page(self, data: Dict[str, Any]):
        """Add cover page - skips if using Word template (template already has cover)"""
        # If using a Word template, the cover page is already there
        # Just add the period information
        if self.using_template:
            self._add_period_info_to_template(data)
            return

        # Fallback to generated cover page
        self._add_generated_cover_page(data)

    def _add_period_info_to_template(self, data: Dict[str, Any]):
        """Add period information when using Word template"""
        period = data.get('period', {})
        start_date = period.get('start', '')
        end_date = period.get('end', '')

        if start_date and end_date:
            period_text = f"POUR LA PERIODE DU {start_date} AU {end_date}"
        else:
            period_text = f"POUR LA PERIODE DU {datetime.now().strftime('%d/%m/%Y')}"

        # Add some spacing after template content
        self.doc.add_paragraph()
        self.doc.add_paragraph()

        # Add period in a bordered box
        title_table = self.doc.add_table(rows=1, cols=1)
        title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        title_cell = title_table.rows[0].cells[0]
        self._set_cell_border(title_cell)

        p_title = title_cell.paragraphs[0]
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run_title = p_title.add_run("RAPPORT DE VEILLE INFORMATIONNELLE SUR\n")
        run_title.bold = True
        run_title.font.size = Pt(14)

        run_title2 = p_title.add_run("L'INTERNET ET LES RESEAUX SOCIAUX\n")
        run_title2.bold = True
        run_title2.font.size = Pt(14)

        run_period = p_title.add_run(period_text)
        run_period.bold = True
        run_period.font.size = Pt(14)

        title_cell.width = Inches(5.5)

    def _add_generated_cover_page(self, data: Dict[str, Any]):
        """Fallback: Generate cover page programmatically"""
        # Create header table for bilingual header (3 columns)
        header_table = self.doc.add_table(rows=2, cols=3)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Set column widths
        for cell in header_table.columns[0].cells:
            cell.width = Inches(2.5)
        for cell in header_table.columns[1].cells:
            cell.width = Inches(1.5)
        for cell in header_table.columns[2].cells:
            cell.width = Inches(2.5)

        # First row - Country names and logo
        cells_row1 = header_table.rows[0].cells

        # Left - French
        p_left = cells_row1[0].paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_left = p_left.add_run("REPUBLIQUE DU CAMEROUN")
        run_left.bold = True
        run_left.font.size = Pt(11)

        # Center - Logo
        p_center = cells_row1[1].paragraphs[0]
        p_center.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._add_logo(p_center)

        # Right - English
        p_right = cells_row1[2].paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_right = p_right.add_run("REPUBLIC OF CAMEROON")
        run_right.bold = True
        run_right.font.size = Pt(11)

        # Second row - Mottos
        cells_row2 = header_table.rows[1].cells

        # Left - French motto
        p_left2 = cells_row2[0].paragraphs[0]
        p_left2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_left2 = p_left2.add_run("Paix – Travail – Patrie")
        run_left2.italic = True
        run_left2.font.size = Pt(10)

        # Center cell (below logo) - empty
        cells_row2[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Right - English motto
        p_right2 = cells_row2[2].paragraphs[0]
        p_right2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_right2 = p_right2.add_run("Peace – Work – Fatherland")
        run_right2.italic = True
        run_right2.font.size = Pt(10)

        self.doc.add_paragraph()  # Spacing

        # Agency names table
        agency_table = self.doc.add_table(rows=1, cols=2)
        agency_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        agency_cells = agency_table.rows[0].cells

        # French agency name
        p_agency_fr = agency_cells[0].paragraphs[0]
        p_agency_fr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_agency_fr = p_agency_fr.add_run("AGENCE NATIONALE DES\nTECHNOLOGIES DE L'INFORMATION\nET DE LA COMMUNICATION")
        run_agency_fr.bold = True
        run_agency_fr.font.size = Pt(10)

        # English agency name
        p_agency_en = agency_cells[1].paragraphs[0]
        p_agency_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_agency_en = p_agency_en.add_run("NATIONAL AGENCY FOR\nINFORMATION AND COMMUNICATION\nTECHNOLOGIES")
        run_agency_en.bold = True
        run_agency_en.font.size = Pt(10)

        # Add spacing to push title box to lower part of page
        for _ in range(12):
            self.doc.add_paragraph()

        # Create bordered title box
        period = data.get('period', {})
        start_date = period.get('start', '')
        end_date = period.get('end', '')

        # Format dates if available
        if start_date and end_date:
            period_text = f"POUR LA PERIODE DU {start_date} AU {end_date}"
        else:
            period_text = f"POUR LA PERIODE DU {datetime.now().strftime('%d/%m/%Y')}"

        # Title box using a single-cell table with border
        title_table = self.doc.add_table(rows=1, cols=1)
        title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        title_cell = title_table.rows[0].cells[0]

        # Set cell borders
        self._set_cell_border(title_cell)

        # Add title text
        p_title = title_cell.paragraphs[0]
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Main title
        run_title = p_title.add_run("RAPPORT DE VEILLE INFORMATIONNELLE SUR\n")
        run_title.bold = True
        run_title.font.size = Pt(14)

        run_title2 = p_title.add_run("L'INTERNET ET LES RESEAUX SOCIAUX\n")
        run_title2.bold = True
        run_title2.font.size = Pt(14)

        run_period = p_title.add_run(period_text)
        run_period.bold = True
        run_period.font.size = Pt(14)

        # Set cell width
        title_cell.width = Inches(5.5)

    def _add_logo(self, paragraph):
        """Add ANTIC logo image or text fallback"""
        # Try to add logo image
        if os.path.exists(self.LOGO_PATH):
            try:
                run = paragraph.add_run()
                run.add_picture(self.LOGO_PATH, width=Inches(1.0))
                return
            except Exception:
                pass

        # Fallback: Create styled text logo
        run_a = paragraph.add_run("A")
        run_a.bold = True
        run_a.font.size = Pt(20)
        run_a.font.color.rgb = RGBColor(0, 128, 0)  # Green

        run_n = paragraph.add_run("N")
        run_n.bold = True
        run_n.font.size = Pt(20)
        run_n.font.color.rgb = RGBColor(255, 204, 0)  # Yellow

        run_t = paragraph.add_run("T")
        run_t.bold = True
        run_t.font.size = Pt(20)
        run_t.font.color.rgb = RGBColor(0, 128, 0)  # Green

        run_i = paragraph.add_run("I")
        run_i.bold = True
        run_i.font.size = Pt(20)
        run_i.font.color.rgb = RGBColor(255, 0, 0)  # Red

        run_c = paragraph.add_run("C")
        run_c.bold = True
        run_c.font.size = Pt(20)
        run_c.font.color.rgb = RGBColor(0, 128, 0)  # Green

    def _set_cell_border(self, cell, border_size=12):
        """Set cell border"""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        for border_name in ['top', 'left', 'bottom', 'right']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), str(border_size))
            border.set(qn('w:color'), '000000')
            tcBorders.append(border)
        tcPr.append(tcBorders)

    def _add_title(self, title: str):
        """Add main title"""
        heading = self.doc.add_heading(title, 0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0, 51, 102)

    def _add_subtitle(self, subtitle: str):
        """Add subtitle"""
        para = self.doc.add_paragraph(subtitle)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.runs[0].font.size = Pt(12)
        para.runs[0].font.italic = True

    def _add_section_heading(self, heading: str):
        """Add section heading"""
        self.doc.add_paragraph()  # Spacing
        h = self.doc.add_heading(heading, level=1)
        for run in h.runs:
            run.font.color.rgb = RGBColor(31, 78, 120)

    def _add_subsection_heading(self, heading: str):
        """Add subsection heading"""
        h = self.doc.add_heading(heading, level=2)
        for run in h.runs:
            run.font.color.rgb = RGBColor(31, 78, 120)
            run.font.size = Pt(12)

    def _add_no_data_message(self, message: str = "Information non disponible pour le moment."):
        """Add a message when no data is available"""
        para = self.doc.add_paragraph()
        run = para.add_run(message)
        run.font.italic = True
        run.font.color.rgb = RGBColor(128, 128, 128)
        run.font.size = Pt(11)

    def _has_data(self, data) -> bool:
        """Check if data has meaningful content"""
        if data is None:
            return False
        if isinstance(data, dict):
            return any(v for v in data.values() if v)
        if isinstance(data, list):
            return len(data) > 0
        return bool(data)

    def _add_presence_passive_table(self, data: Dict[str, Any]):
        """Add presence passive indicators table"""
        passive = data.get('presence_passive', {})

        # Check if we have any data
        if not self._has_data(passive):
            self._add_no_data_message("Les indicateurs de présence passive ne sont pas disponibles pour ce rapport.")
            return

        table = self.doc.add_table(rows=4, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header row
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Indicateur'
        header_cells[1].text = 'Valeur'
        header_cells[2].text = 'Évolution'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        # Followers row
        cells1 = table.rows[1].cells
        cells1[0].text = 'Nombre de Followers'
        cells1[1].text = f"{passive.get('followers', 0):,}"
        evolution = passive.get('followers_evolution', 0)
        cells1[2].text = f"{evolution:+.2f}%" if evolution else "N/A"

        # Views/Impressions row
        cells2 = table.rows[2].cells
        cells2[0].text = 'Nombre de Vues/Impressions'
        cells2[1].text = f"{passive.get('views', 0):,}"
        views_evolution = passive.get('views_evolution', 0)
        cells2[2].text = f"{views_evolution:+.2f}%" if views_evolution else "N/A"

        # Potential Reach row
        cells3 = table.rows[3].cells
        cells3[0].text = 'Portée Potentielle (Potential Reach)'
        cells3[1].text = f"{passive.get('potential_reach', 0):,}"
        reach_evolution = passive.get('reach_evolution', 0)
        cells3[2].text = f"{reach_evolution:+.2f}%" if reach_evolution else "N/A"

    def _add_presence_active_table(self, data: Dict[str, Any]):
        """Add presence active indicators table"""
        active = data.get('presence_active', {})

        # Check if we have any data
        if not self._has_data(active):
            self._add_no_data_message("Les indicateurs de présence active ne sont pas disponibles pour ce rapport.")
            return

        table = self.doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'

        # Header row
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Indicateur'
        header_cells[1].text = 'Valeur'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')

        # Comments row
        cells1 = table.rows[1].cells
        cells1[0].text = 'Nombre de Commentaires'
        cells1[1].text = f"{active.get('comments', 0):,}"

        # Likes row
        cells2 = table.rows[2].cells
        cells2[0].text = "Nombre de Likes (J'aime)"
        cells2[1].text = f"{active.get('likes', 0):,}"

        # Shares row
        cells3 = table.rows[3].cells
        cells3[0].text = 'Nombre de Partages (Shares)'
        cells3[1].text = f"{active.get('shares', 0):,}"

    def _add_opinion_trends_table(self, sentiment_data: Dict[str, int]):
        """Add opinion trends table"""
        # Check if we have any data
        if not self._has_data(sentiment_data):
            self._add_no_data_message("Les indicateurs des tendances d'opinions ne sont pas disponibles pour ce rapport.")
            return

        table = self.doc.add_table(rows=4, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Tendance'
        header_cells[1].text = 'Nombre'
        header_cells[2].text = 'Pourcentage'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        total = sum(sentiment_data.values()) if sentiment_data else 0

        # Positive row
        cells1 = table.rows[1].cells
        positive = sentiment_data.get('positive', 0)
        cells1[0].text = 'Positif'
        cells1[1].text = str(positive)
        cells1[2].text = f"{(positive / total * 100):.2f}%" if total > 0 else "0%"

        # Neutral row
        cells2 = table.rows[2].cells
        neutral = sentiment_data.get('neutral', 0)
        cells2[0].text = 'Neutre'
        cells2[1].text = str(neutral)
        cells2[2].text = f"{(neutral / total * 100):.2f}%" if total > 0 else "0%"

        # Negative row
        cells3 = table.rows[3].cells
        negative = sentiment_data.get('negative', 0)
        cells3[0].text = 'Négatif'
        cells3[1].text = str(negative)
        cells3[2].text = f"{(negative / total * 100):.2f}%" if total > 0 else "0%"

    def _add_sources_table(self, sources_data: Dict[str, int]):
        """Add sources distribution table"""
        if not self._has_data(sources_data):
            self._add_no_data_message("La répartition par source n'est pas disponible pour ce rapport.")
            return

        total = sum(sources_data.values())

        # Sort by count descending
        sorted_sources = sorted(sources_data.items(), key=lambda x: x[1], reverse=True)

        table = self.doc.add_table(rows=len(sorted_sources) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Plateforme'
        header_cells[1].text = 'Mentions'
        header_cells[2].text = 'Pourcentage'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        # Data rows
        for idx, (source, count) in enumerate(sorted_sources, 1):
            cells = table.rows[idx].cells
            cells[0].text = source
            cells[1].text = str(count)
            percentage = (count / total * 100) if total > 0 else 0
            cells[2].text = f"{percentage:.2f}%"

    def _add_languages_table(self, languages_data: Dict[str, int]):
        """Add languages distribution table"""
        if not self._has_data(languages_data):
            self._add_no_data_message("La répartition par langue n'est pas disponible pour ce rapport.")
            return

        total = sum(languages_data.values())

        table = self.doc.add_table(rows=len(languages_data) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Langue'
        header_cells[1].text = 'Mentions'
        header_cells[2].text = 'Pourcentage'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        # Data rows
        for idx, (language, count) in enumerate(languages_data.items(), 1):
            cells = table.rows[idx].cells
            cells[0].text = language
            cells[1].text = str(count)
            percentage = (count / total * 100) if total > 0 else 0
            cells[2].text = f"{percentage:.2f}%"

    def _add_topics_table(self, topics: list):
        """Add top topics table"""
        if not self._has_data(topics):
            self._add_no_data_message("Les sujets principaux ne sont pas disponibles pour ce rapport.")
            return

        # Limit to top 15
        topics = topics[:15]

        table = self.doc.add_table(rows=len(topics) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Rang'
        header_cells[1].text = 'Sujet'
        header_cells[2].text = 'Mentions'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        # Data rows
        for idx, topic in enumerate(topics, 1):
            cells = table.rows[idx].cells
            cells[0].text = str(idx)
            cells[1].text = topic.get('name', '')
            cells[2].text = str(topic.get('count', 0))

    def _add_hashtags_table(self, hashtags: list):
        """Add hashtags table"""
        if not self._has_data(hashtags):
            self._add_no_data_message("Les hashtags populaires ne sont pas disponibles pour ce rapport.")
            return

        # Limit to top 10
        hashtags = hashtags[:10]

        table = self.doc.add_table(rows=len(hashtags) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Rang'
        header_cells[1].text = 'Hashtag'
        header_cells[2].text = 'Occurrences'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        # Data rows
        for idx, hashtag in enumerate(hashtags, 1):
            cells = table.rows[idx].cells
            cells[0].text = str(idx)
            cells[1].text = hashtag.get('hashtag', '')
            cells[2].text = str(hashtag.get('count', 0))

    def _add_emotions_table(self, emotion_data: Dict[str, int]):
        """Add emotions distribution table"""
        if not self._has_data(emotion_data):
            self._add_no_data_message("Les indicateurs des émotions ne sont pas disponibles pour ce rapport.")
            return

        # Define emotion labels in French
        emotion_labels = {
            'joy': 'Joie',
            'neutral': 'Neutre',
            'anger': 'Colère',
            'sadness': 'Tristesse',
            'surprise': 'Surprise',
            'fear': 'Peur',
            'disgust': 'Dégoût'
        }

        total = sum(emotion_data.values())

        # Sort emotions by count
        sorted_emotions = sorted(emotion_data.items(), key=lambda x: x[1], reverse=True)

        table = self.doc.add_table(rows=len(sorted_emotions) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Émotion'
        header_cells[1].text = 'Nombre'
        header_cells[2].text = 'Pourcentage'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        # Data rows
        for idx, (emotion, count) in enumerate(sorted_emotions, 1):
            cells = table.rows[idx].cells
            cells[0].text = emotion_labels.get(emotion, emotion.capitalize())
            cells[1].text = str(count)
            percentage = (count / total * 100) if total > 0 else 0
            cells[2].text = f"{percentage:.2f}%"

    def _add_influencers_table(self, influencers: list):
        """Add influencers table"""
        if not self._has_data(influencers):
            self._add_no_data_message("Les influenceurs ne sont pas disponibles pour ce rapport.")
            return

        # Limit to top 10
        influencers = influencers[:10]

        table = self.doc.add_table(rows=len(influencers) + 1, cols=4)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Rang'
        header_cells[1].text = 'Nom'
        header_cells[2].text = 'Plateforme'
        header_cells[3].text = 'Score d\'influence'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')
        self._set_cell_background(header_cells[3], 'D5E8F0')

        # Data rows
        for idx, influencer in enumerate(influencers, 1):
            cells = table.rows[idx].cells
            cells[0].text = str(idx)
            cells[1].text = influencer.get('name', '')
            cells[2].text = influencer.get('platform', 'Facebook')
            score = influencer.get('influence_score', 0)
            # Show "-" if no score available
            cells[3].text = f"{score}/100" if score > 0 else "-"

    def _add_reach_breakdown_table(self, reach_posts: list):
        """Add reach breakdown (top posts) table"""
        if not self._has_data(reach_posts):
            self._add_no_data_message("La répartition de la portée n'est pas disponible pour ce rapport.")
            return

        # Limit to top 5
        reach_posts = reach_posts[:5]

        table = self.doc.add_table(rows=len(reach_posts) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Rang'
        header_cells[1].text = 'Auteur/Source'
        header_cells[2].text = 'Portée'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        # Data rows
        for idx, post in enumerate(reach_posts, 1):
            cells = table.rows[idx].cells
            cells[0].text = str(idx)
            cells[1].text = post.get('name', '')
            reach_display = post.get('reach_display', '')
            reach_num = post.get('reach', 0)
            cells[2].text = reach_display if reach_display else f"{reach_num:,}"

    # Brand24-specific methods
    def _add_presence_score_info(self, presence_score: Dict[str, Any]):
        """Add presence score information (Brand24)"""
        if not self._has_data(presence_score):
            self._add_no_data_message("Le score de présence n'est pas disponible.")
            return

        table = self.doc.add_table(rows=2, cols=2)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Indicateur'
        header_cells[1].text = 'Valeur'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')

        cells = table.rows[1].cells
        cells[0].text = 'Score de Présence'
        score = presence_score.get('score', 0)
        percentile = presence_score.get('percentile', 0)
        cells[1].text = f"{score}/100 (supérieur à {percentile}% des marques)"

    def _add_ave_info(self, ave: Dict[str, Any]):
        """Add AVE (Advertising Value Equivalent) information"""
        if not self._has_data(ave):
            self._add_no_data_message("La valeur publicitaire équivalente n'est pas disponible.")
            return

        para = self.doc.add_paragraph()
        value = ave.get('value', 0)
        currency = ave.get('currency', 'USD')
        run = para.add_run(f"Valeur estimée: ${value:,} {currency}")
        run.font.size = Pt(12)
        run.bold = True

    # Demographics-specific methods
    def _add_reach_summary(self, reach: Dict[str, Any]):
        """Add reach summary for demographics report"""
        if not self._has_data(reach):
            self._add_no_data_message("Les données de portée ne sont pas disponibles.")
            return

        para = self.doc.add_paragraph()
        reach_value = reach.get('reach', 0)
        run = para.add_run(f"Portée totale: {reach_value:,}")
        run.font.size = Pt(14)
        run.bold = True

    def _add_gender_table(self, gender: Dict[str, float]):
        """Add gender distribution table"""
        if not self._has_data(gender):
            self._add_no_data_message("La répartition par genre n'est pas disponible.")
            return

        table = self.doc.add_table(rows=3, cols=2)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Genre'
        header_cells[1].text = 'Pourcentage'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')

        # Data
        cells1 = table.rows[1].cells
        cells1[0].text = 'Femme'
        cells1[1].text = f"{gender.get('female', 0):.1f}%"

        cells2 = table.rows[2].cells
        cells2[0].text = 'Homme'
        cells2[1].text = f"{gender.get('male', 0):.1f}%"

    def _add_age_table(self, age_data: list):
        """Add age distribution table"""
        if not self._has_data(age_data):
            self._add_no_data_message("La répartition par âge n'est pas disponible.")
            return

        table = self.doc.add_table(rows=len(age_data) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Tranche d\'âge'
        header_cells[1].text = 'Pourcentage'
        header_cells[2].text = 'Nombre'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        for idx, item in enumerate(age_data, 1):
            cells = table.rows[idx].cells
            cells[0].text = item.get('age_group', '')
            cells[1].text = f"{item.get('percentage', 0):.1f}%"
            cells[2].text = f"{item.get('count', 0):,}"

    def _add_countries_table(self, countries: list):
        """Add countries distribution table"""
        if not self._has_data(countries):
            self._add_no_data_message("La répartition par pays n'est pas disponible.")
            return

        table = self.doc.add_table(rows=len(countries) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Pays'
        header_cells[1].text = 'Pourcentage'
        header_cells[2].text = 'Portée'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        for idx, item in enumerate(countries, 1):
            cells = table.rows[idx].cells
            cells[0].text = item.get('country', '')
            cells[1].text = f"{item.get('percentage', 0):.1f}%"
            cells[2].text = f"{item.get('reach', 0):,}"

    def _add_occupation_table(self, occupations: list):
        """Add occupation distribution table"""
        if not self._has_data(occupations):
            self._add_no_data_message("La répartition par occupation n'est pas disponible.")
            return

        table = self.doc.add_table(rows=len(occupations) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Occupation'
        header_cells[1].text = 'Pourcentage'
        header_cells[2].text = 'Nombre'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        for idx, item in enumerate(occupations, 1):
            cells = table.rows[idx].cells
            cells[0].text = item.get('occupation', '')
            cells[1].text = f"{item.get('percentage', 0):.1f}%"
            cells[2].text = f"{item.get('count', 0):,}"

    def _add_education_table(self, education: list):
        """Add education distribution table"""
        if not self._has_data(education):
            self._add_no_data_message("La répartition par niveau d'éducation n'est pas disponible.")
            return

        table = self.doc.add_table(rows=len(education) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Niveau d\'éducation'
        header_cells[1].text = 'Pourcentage'
        header_cells[2].text = 'Nombre'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        for idx, item in enumerate(education, 1):
            cells = table.rows[idx].cells
            cells[0].text = item.get('level', '')
            cells[1].text = f"{item.get('percentage', 0):.1f}%"
            cells[2].text = f"{item.get('count', 0):,}"

    def _add_interests_table(self, interests: list):
        """Add interests distribution table"""
        if not self._has_data(interests):
            self._add_no_data_message("Les centres d'intérêt ne sont pas disponibles.")
            return

        table = self.doc.add_table(rows=len(interests) + 1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Centre d\'intérêt'
        header_cells[1].text = 'Pourcentage'
        header_cells[2].text = 'Portée'
        self._set_cell_background(header_cells[0], 'D5E8F0')
        self._set_cell_background(header_cells[1], 'D5E8F0')
        self._set_cell_background(header_cells[2], 'D5E8F0')

        for idx, item in enumerate(interests, 1):
            cells = table.rows[idx].cells
            cells[0].text = item.get('interest', '')
            cells[1].text = f"{item.get('percentage', 0):.1f}%"
            cells[2].text = f"{item.get('reach', 0):,}"

    def _set_cell_background(self, cell, color: str):
        """Set cell background color"""
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), color)
        cell._element.get_or_add_tcPr().append(shading_elm)


def generate_word_report(data: Dict[str, Any], output_path: str) -> str:
    """Convenience function to generate Word report"""
    generator = WordReportGenerator()
    return generator.generate_report(data, output_path)
